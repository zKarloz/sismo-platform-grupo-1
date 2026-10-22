"""Lectura en modo solo lectura del catálogo instrumental IGP exportado a XLSX."""

from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

# =====================================================
# BLOQUE 1: Contrato de columnas de la exportación IGP.
# La cabecera es parte del contrato: un cambio exige
# revisión explícita, no mapeos silenciosos.
# =====================================================

SHEET_NAME = "IGP - Datos sísmicos"
SOURCE_FIELDS = (
    "fecha UTC",
    "hora UTC",
    "latitud (º)",
    "longitud (º)",
    "profundidad (km)",
    "magnitud (M)",
)


# =====================================================
# BLOQUE 2: Lector iterativo y conservador del archivo.
# No guarda cambios, no elimina registros y devuelve
# el número físico de la fila para trazabilidad.
# =====================================================

def iter_igp_rows(path: str | Path) -> Iterator[tuple[int, dict]]:
    input_path = Path(path)
    if not input_path.is_file():
        raise FileNotFoundError(f"No existe el Excel IGP: {input_path}")

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"No existe la hoja esperada '{SHEET_NAME}'. "
                f"Hojas disponibles: {workbook.sheetnames}"
            )

        worksheet = workbook[SHEET_NAME]
        iterator = worksheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            raise ValueError("El Excel IGP no contiene una cabecera")

        headers = [str(h).strip() if h is not None else "" for h in header_row]
        if len(headers) != len(SOURCE_FIELDS) or set(headers) != set(SOURCE_FIELDS):
            raise ValueError(
                "Cabecera IGP inesperada; revise manualmente la versión del archivo. "
                f"Esperadas: {list(SOURCE_FIELDS)}; encontradas: {headers}"
            )

        for row_number, cells in enumerate(iterator, start=2):
            raw = dict(zip(headers, cells))
            # Conservar incluso filas con valores nulos para registrar errores.
            yield row_number, raw
    finally:
        workbook.close()
